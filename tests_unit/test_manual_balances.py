from copy import deepcopy
from io import BytesIO
import unittest

from openpyxl import load_workbook

from core.calculation import build_snapshot
from core.export import export_xlsx
from test_accounting import accounting_fixture
from test_calculation import fixture, ledger, movement
from test_template import row_for, saved, template_workbook


def capture(opening, closing, ident=1):
    return {"source_id": ident, "source_model": "cq.bank.balance", "name": "Estado de prueba",
            "opening": opening, "closing": closing, "currency": "GTQ", "note": "Revisado",
            "write_uid": 2, "write_date": "2024-04-01 12:00:00"}


class TestManualBalances(unittest.TestCase):
    def test_monthly_balances_preserve_books_and_do_not_invent_bank_flows(self):
        data = accounting_fixture()
        data["manual_bank_balances"] = {"1": capture(1000, 1195), "2": capture(1195, 1170, 2)}
        before = deepcopy(data)
        result = build_snapshot(data)
        jan, feb, mar = result["months"]
        self.assertEqual([month["reported_bank_opening"] for month in result["months"]], [1000, 1195, 1170])
        self.assertEqual([month["statement_end"] for month in result["months"]], [1195, 1170, None])
        self.assertTrue(all(month["bank_end"] is None for month in result["months"]))
        self.assertEqual((jan["bank_adjusted"], feb["bank_adjusted"], mar["bank_adjusted"]), (1195, 1170, None))
        self.assertEqual((jan["difference"], feb["difference"]), (0, 0))
        self.assertEqual(result["accounting_months"][0]["income"], 400)
        self.assertFalse(result["movements"])
        self.assertEqual({(issue["code"], issue.get("month")) for issue in result["issues"]},
                         {("bank_flow_missing", 1), ("bank_flow_missing", 2), ("coverage", 3)})
        self.assertEqual(data, before)
        data["manual_bank_balances"]["1"]["closing"] = 9999
        self.assertEqual(jan["manual_balance"]["closing"], 1195)

    def test_real_zero_is_reported_but_does_not_prove_zero_flows(self):
        data = fixture()
        data.update(bank_opening=None, movements=[], ledger=[], controls={}, manual_bank_balances={"1": capture(0, 0)})
        result = build_snapshot(data)
        jan = result["months"][0]
        self.assertEqual((jan["reported_bank_opening"], jan["statement_end"], jan["difference"]), (0, 0, 0))
        self.assertIsNone(jan["bank_end"])
        self.assertEqual([issue["code"] for issue in result["issues"]], ["bank_flow_missing"])

    def test_native_statements_win_and_manual_discrepancies_are_visible(self):
        data = fixture()
        data["manual_bank_balances"] = {"1": capture(1100, 1250)}
        result = build_snapshot(data)
        jan = result["months"][0]
        self.assertEqual((jan["reported_bank_opening"], jan["statement_end"], jan["bank_end"]), (1000, 1195, 1195))
        self.assertEqual(jan["control"]["name"], "ENERO")
        self.assertEqual({issue["code"] for issue in result["issues"]}, {"manual_bank_opening_conflict", "manual_bank_closing_conflict"})

    def test_manual_opening_can_anchor_real_bank_transactions(self):
        data = fixture()
        data.update(month=2, bank_opening=None, controls={}, manual_bank_balances={"2": capture(1195, 1170)})
        data["movements"].append(movement(3, "2024-02-10", -25, "OUT_SUPPLIERS"))
        data["ledger"].append(ledger(13, "2024-02-10", -25))
        result = build_snapshot(data)
        self.assertEqual([month["bank_opening"] for month in result["months"]], [1000, 1195])
        self.assertEqual([month["bank_end"] for month in result["months"]], [1195, 1170])
        self.assertEqual([row["running_balance"] for row in result["movements"]], [1400, 1195, 1170])
        self.assertEqual(result["issues"][0]["code"], "coverage")
        self.assertEqual(len(result["issues"]), 1)  # January closing is still absent.

    def test_continuity_and_closing_differences_remain_review_items(self):
        data = accounting_fixture()
        data["manual_bank_balances"] = {"1": capture(1000, 1195), "2": capture(1200, 1190)}
        result = build_snapshot(data)
        feb = result["months"][1]
        self.assertEqual(feb["difference"], 20)
        codes = {item["code"] for item in result["issues"] if item.get("month") == 2}
        self.assertTrue({"bank_flow_missing", "manual_bank_continuity", "book_difference"} <= codes)

    def test_declared_closing_applies_pending_items_once(self):
        data = accounting_fixture()
        data.update(month=1, manual_bank_balances={"1": capture(1000, 1195)})
        data["ledger"].extend([ledger(30, "2024-01-20", 50, "outstanding"),
                               ledger(31, "2024-01-20", -20, "outstanding", method="check_printing"),
                               ledger(32, "2024-01-20", -10, "outstanding")])
        jan = build_snapshot(data)["months"][0]
        self.assertEqual((jan["bank_adjusted"], jan["book_adjusted"], jan["difference"]), (1215, 1215, 0))

    def test_exports_fill_captures_and_keep_missing_flows_unknown(self):
        data = accounting_fixture()
        data["manual_bank_balances"] = {"1": capture(1000, 1195), "2": capture(1195, 1170)}
        template = template_workbook()
        template["Banco"]["C130"] = "Banco inicial"
        template["Banco"]["G130"] = 999999
        for source in (None, saved(template)):
            with self.subTest(template=bool(source)):
                output = export_xlsx(build_snapshot(data), template=source)
                book = load_workbook(BytesIO(output), data_only=True)
                formulas = load_workbook(BytesIO(output))["Banco"]
                bank = book["Banco"]
                self.assertEqual(book.active.title, "Resumen contable Odoo")
                self.assertEqual([bank.cell(19, c).value for c in (7, 9, 11, 31)], [1000, 1195, 1170, 1000])
                self.assertEqual(bank["G20"].value, "n.d.")
                self.assertEqual(bank["G21"].value, "n.d.")
                self.assertIsNone(bank["M19"].value)
                self.assertEqual(book["Data"]["K8"].value, 1000)
                closing_row = row_for(bank, "Saldo final según banco (extracto o captura)")
                adjusted_row = row_for(bank, "SALDO CONCILIADO")
                self.assertEqual(bank.cell(closing_row, 7).value, 1195)
                self.assertEqual(bank.cell(adjusted_row, 7).value, 1195)
                self.assertEqual(bank.cell(adjusted_row, 11).value, "n.d.")
                self.assertIn("IF(ISNUMBER(", formulas.cell(adjusted_row, 7).value)
                self.assertFalse(any(c.data_type == "e" for s in book for row in s for c in row))
                if source:
                    footer = row_for(bank, "Banco inicial")
                    self.assertEqual([bank.cell(footer, c).value for c in (7, 9, 11, 31)], [1000, 1195, 1170, 1000])
