from io import BytesIO
from pathlib import Path
import unittest
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.styles import PatternFill

from core.calculation import build_snapshot
from core.export import export_xlsx
from core.template import read_template
from test_calculation import fixture, ledger, movement


def full_fixture():
    data = fixture()
    concepts = ET.parse(Path(__file__).parents[1] / "data/concepts.xml").getroot()
    data["concepts"] = [{field.attrib["name"]: field.text for field in record} for record in concepts]
    return data


def template_workbook():
    raw = export_xlsx(build_snapshot(full_fixture()))
    end = read_template(raw)["form_end"]
    book = openpyxl.load_workbook(BytesIO(raw))
    for merged in list(book["Banco"].merged_cells.ranges):
        if merged.min_row > end:
            book["Banco"].unmerge_cells(str(merged))
    book["Banco"].delete_rows(end + 1, book["Banco"].max_row)
    del book["Partidas conciliatorias"]
    return book


def saved(book):
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def row_for(sheet, label, last=False):
    rows = [row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 3).value == label]
    return rows[-1 if last else 0]


class TestTemplate(unittest.TestCase):
    def test_usd_variant_keeps_data_header_on_row_eight(self):
        template = template_workbook()
        template["Data"].insert_rows(7)
        data = fixture()
        data["metadata"]["currency"] = "USD"
        result = export_xlsx(build_snapshot(data), template=saved(template))
        book = openpyxl.load_workbook(BytesIO(result), data_only=True)
        self.assertEqual(book["Data"]["A8"].value, "Mes de cobro")
        self.assertIsNone(book["Data"]["A7"].value)
        self.assertEqual(book["Data"]["K9"].value, 1000)
        self.assertEqual(book["Data"]["K11"].value, 1195)
        self.assertEqual(book["Banco"]["F8"].value, "Dólares")
        self.assertEqual(book["Banco"]["F35"].value, "Valor $")
    def test_all_twelve_month_openings_and_empty_months(self):
        data = fixture()
        data["month"] = 12
        for month in (2, 4, 8, 12):
            data["movements"].append(movement(month + 100, "2024-%02d-10" % month, -month, "OUT_FEES"))
        snapshot = build_snapshot(data)
        output = export_xlsx(snapshot, template=saved(template_workbook()))
        values = openpyxl.load_workbook(BytesIO(output), data_only=True)["Banco"]
        formulas = openpyxl.load_workbook(BytesIO(output))["Banco"]
        ending = row_for(values, "SALDO FINAL BANCARIO ( A - B )")
        for idx, month in enumerate(snapshot["months"]):
            col = 7 + idx * 2
            self.assertEqual(values.cell(19, col).value, month["bank_opening"])
            if idx:
                self.assertEqual(formulas.cell(19, col).value, "=%s%s" % (openpyxl.utils.get_column_letter(col - 2), ending))
        self.assertEqual(values["AE19"].value, 1000)

    def test_manual_notes_style_and_old_amount_replacement(self):
        template = template_workbook()
        main = template["Banco"]
        main["G19"], main["I19"], main["G21"] = 999999, 777777, 555555
        main["D6"] = "Banco escrito a mano"
        main["G19"].fill = PatternFill("solid", fgColor="FFF2CC")
        main.merge_cells("H4:K4")
        main["H4"] = "Preparó: persona de prueba"
        main["C120"] = "Firma y observación manual"
        template.create_sheet("Notas manuales")["B1"] = 123.45
        template["Data"]["I400"] = 888888
        data = fixture()
        data["metadata"]["bank"] = ""
        output = export_xlsx(build_snapshot(data), template=saved(template))
        values = openpyxl.load_workbook(BytesIO(output), data_only=True)
        main = values["Banco"]
        self.assertEqual(main["G19"].value, 1000)
        self.assertIsNone(main["I19"].value)
        self.assertEqual(main["G21"].value, 400)
        self.assertEqual(main["G19"].fill.fgColor.rgb, "00FFF2CC")
        self.assertEqual(main["D6"].value, "Banco escrito a mano")
        self.assertEqual(main["H4"].value, "Preparó: persona de prueba")
        self.assertIn("H4:K4", {str(merged) for merged in main.merged_cells.ranges})
        self.assertEqual(main["C120"].value, "Firma y observación manual")
        self.assertEqual(values["Notas manuales"]["B1"].value, 123.45)
        self.assertIsNone(values["Data"]["I400"].value)
        self.assertEqual(values["Data"]["K10"].value, 1195)
        self.assertIn("Control Odoo", values.sheetnames)

    def test_manual_bank_matching_and_detail_expansion(self):
        template = template_workbook()
        template["Banco"]["C36"], template["Banco"]["D36"] = "Banco manual conocido", "000555"
        template["Banco"]["C37"], template["Banco"]["D37"] = "Banco sin correspondencia", "000999"
        old_ending = row_for(template["Banco"], "SALDO FINAL BANCARIO ( A - B )")
        template.create_sheet("Notas")["A1"] = "='Banco'!$G$%s" % old_ending
        data = full_fixture()
        data["movements"] = []
        for idx in range(8):
            item = movement(idx + 1, "2024-01-10", 10, "IN_TRANSFER")
            item["counterparty_account"] = "000555" if idx == 0 else "000%s" % idx
            item["counterparty_bank"] = "" if idx == 0 else "Banco de prueba %s" % idx
            data["movements"].append(item)
        output = export_xlsx(build_snapshot(data), template=saved(template))
        values = openpyxl.load_workbook(BytesIO(output), data_only=True)["Banco"]
        formulas = openpyxl.load_workbook(BytesIO(output))
        self.assertEqual(values["C36"].value, "Banco manual conocido")
        self.assertEqual(values["F36"].value, 10)
        self.assertEqual(values["C37"].value, "Banco sin correspondencia")
        self.assertEqual(values["F37"].value, 0)
        self.assertEqual(values["G34"].value, 80)
        ending = row_for(values, "SALDO FINAL BANCARIO ( A - B )")
        self.assertGreater(ending, old_ending)
        self.assertEqual(values.cell(ending, 7).value, 1080)
        self.assertEqual(formulas["Notas"]["A1"].value, "='Banco'!$G$%s" % ending)

    def test_custom_and_unclassified_flows_count_once(self):
        data = full_fixture()
        data["concepts"].append({"code": "CUSTOM", "name": "Ingreso especial", "direction": "in", "detail": "none"})
        data["movements"] = [movement(1, "2024-01-10", 100, "CUSTOM"), movement(2, "2024-01-15", -80, allocations=[])]
        sheet = openpyxl.load_workbook(BytesIO(export_xlsx(build_snapshot(data))), data_only=True)["Banco"]
        self.assertEqual(sheet.cell(row_for(sheet, "Otros ingresos (especifique)"), 7).value, 100)
        self.assertEqual(sheet.cell(row_for(sheet, "Otros egresos (especifique):"), 7).value, 80)
        self.assertEqual(sheet.cell(row_for(sheet, "TOTAL EGRESOS"), 31).value, 80)
        self.assertEqual(sheet.cell(row_for(sheet, "SALDO FINAL BANCARIO ( A - B )"), 31).value, 1020)

    def test_books_remain_independent_and_interest_is_not_adjusted_twice(self):
        data = full_fixture()
        data["movements"][0]["allocations"][0]["code"] = "IN_INTEREST"
        data["ledger"].append(ledger(13, "2024-01-10", -5))
        output = export_xlsx(build_snapshot(data), template=saved(template_workbook()))
        sheet = openpyxl.load_workbook(BytesIO(output), data_only=True)["Banco"]
        self.assertEqual(sheet.cell(row_for(sheet, "Ingresos por intereses ganados"), 7).value, 0)
        self.assertEqual(sheet.cell(row_for(sheet, "SALDO CONCILIADO"), 7).value, 1195)
        self.assertEqual(sheet.cell(row_for(sheet, "SALDO CONCILIADO", last=True), 7).value, 1190)

    def test_text_safety_and_account_numbers_survive_template(self):
        data = fixture()
        data["movements"][0]["description"] = '=HYPERLINK("https://invalid.test","x")'
        data["movements"][0]["counterparty_account"] = "000123"
        workbook = openpyxl.load_workbook(BytesIO(export_xlsx(build_snapshot(data), template=saved(template_workbook()))))
        self.assertEqual(workbook["Data"]["G9"].data_type, "s")
        self.assertEqual(workbook["Data"]["S9"].value, "000123")
        self.assertEqual(workbook["Banco"]["D8"].value, "0000123")
        self.assertEqual(workbook["Banco"]["D8"].number_format, "@")

    def test_incompatible_template_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "XLSX válida"):
            read_template(b"not a workbook")
        template = template_workbook()
        template["Banco"]["F18"] = "Otro encabezado"
        with self.assertRaisesRegex(ValueError, "enero a diciembre"):
            read_template(saved(template))
