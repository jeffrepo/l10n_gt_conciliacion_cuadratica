from io import BytesIO
import unittest
from zipfile import ZipFile

import openpyxl

from core.calculation import build_snapshot
from core.export import export_xlsx, safe_filename
from test_calculation import fixture, ledger, movement


class TestExport(unittest.TestCase):
    def test_annual_totals_and_opening_roll_forward(self):
        data = fixture()
        data["month"] = 2
        data["movements"].append(movement(3, "2024-02-10", -25, "OUT_FEES"))
        data["ledger"].append(ledger(13, "2024-02-10", -25))
        data["controls"]["2"] = {"name": "FEB", "source_id": 2, "date": "2024-02-29", "balance": 1170, "complete": True}
        output = export_xlsx(build_snapshot(data))
        values = openpyxl.load_workbook(BytesIO(output), data_only=True)
        formulas = openpyxl.load_workbook(BytesIO(output), data_only=False)
        self.assertEqual(values.sheetnames, ["Conciliación", "Movimientos", "Partidas conciliatorias"])
        sheet = values["Conciliación"]
        rows = {sheet.cell(row, 2).value: row for row in range(1, sheet.max_row + 1)}
        self.assertEqual(sheet.cell(rows["Total egresos del período"], 15).value, 230)
        self.assertEqual(sheet.cell(rows["Saldo final bancario calculado"], 15).value, 1170)
        self.assertEqual(sheet.cell(rows["Saldo inicial según banco"], 15).value, 1000)
        self.assertEqual(sheet.cell(rows["Saldo inicial según banco"], 4).value, 1195)
        self.assertIsNone(sheet.cell(rows["Saldo inicial según banco"], 5).value)
        self.assertTrue(formulas["Conciliación"].cell(rows["Total egresos del período"], 15).value.startswith("=SUM("))
        self.assertEqual(formulas["Conciliación"].cell(rows["Saldo inicial según banco"], 4).value, "=C%s" % rows["Saldo final bancario calculado"])
        for worksheet in values:
            self.assertFalse(any(cell.data_type == "e" for row in worksheet for cell in row))

    def test_user_text_cannot_become_formula_and_account_keeps_zeros(self):
        data = fixture()
        data["movements"][0]["description"] = '=HYPERLINK("https://invalid.test","x")'
        data["movements"][0]["counterparty_account"] = "000123"
        workbook = openpyxl.load_workbook(BytesIO(export_xlsx(build_snapshot(data))), data_only=False)
        cell = workbook["Movimientos"]["F2"]
        self.assertEqual(cell.data_type, "s")
        self.assertEqual(workbook["Movimientos"]["Q2"].value, "000123")

    def test_split_counts_once_and_preserves_company_total(self):
        data = fixture()
        data["movements"][1]["allocations"] = [{"code": "OUT_SUPPLIERS", "amount": 200, "origin": "manual"}, {"code": "OUT_FEES", "amount": 5, "origin": "manual"}]
        data["movements"][1]["company_amount"] = -1599
        book = openpyxl.load_workbook(BytesIO(export_xlsx(build_snapshot(data))), data_only=True)
        sheet = book["Movimientos"]
        self.assertEqual(sum(sheet.cell(row, 11).value or 0 for row in range(2, sheet.max_row + 1)), 205)
        self.assertAlmostEqual(sheet["M3"].value + sheet["M4"].value, -1599)
        self.assertIsNone(sheet["L3"].value)
        self.assertEqual(sheet["L4"].value, 1195)

    def test_missing_bank_data_is_visible(self):
        data = fixture()
        data["bank_opening"] = None
        book = openpyxl.load_workbook(BytesIO(export_xlsx(build_snapshot(data))), data_only=True)
        self.assertEqual(book["Conciliación"]["C11"].value, "n.d.")
        self.assertIn("BORRADOR", book["Conciliación"]["A7"].value)

    def test_no_external_workbook_links_and_safe_filename(self):
        result = export_xlsx(build_snapshot(fixture()))
        with ZipFile(BytesIO(result)) as archive:
            self.assertFalse(any("externalLinks" in name for name in archive.namelist()))
        self.assertEqual(safe_filename("BAC $ / Cuenta Q"), "BAC_Cuenta_Q")


if __name__ == "__main__":
    unittest.main()
