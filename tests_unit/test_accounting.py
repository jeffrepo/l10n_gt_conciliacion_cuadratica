from io import BytesIO
import unittest

from openpyxl import load_workbook

from core.calculation import build_snapshot
from core.export import export_xlsx
from test_calculation import fixture, ledger, movement
from test_template import saved, template_workbook


def accounting_fixture():
    data = fixture()
    data.update(month=3, bank_opening=None, movements=[], controls={})
    data["ledger"].append(ledger(13, "2024-02-10", -25))
    data["accounting_movements"] = [
        {**row, "journal": "Banco", "reference": "=Not a formula", "concept_code": code,
         "classification_origin": "reconciliation" if code else "pending",
         "counterpart_accounts": "Clientes", "statement_line_id": False}
        for row, code in zip(data["ledger"][1:], ["IN_CUSTOMERS_LOCAL", "OUT_SUPPLIERS", ""])
    ]
    return data


class TestAccounting(unittest.TestCase):
    def test_books_available_without_fabricating_bank_movements_or_opening(self):
        data = accounting_fixture()
        result = build_snapshot(data)
        self.assertEqual(result["movements"], [])
        self.assertEqual(result["months"][0]["income"], 0)
        self.assertIsNone(result["months"][0]["bank_opening"])
        months = result["accounting_months"]
        self.assertEqual((months[0]["opening"], months[0]["income"], months[0]["expense"], months[0]["closing"]), (1000, 400, 205, 1195))
        self.assertEqual((months[1]["opening"], months[1]["closing"], months[2]["closing"]), (1195, 1170, 1170))
        self.assertEqual(months[0]["concept_totals"]["IN_CUSTOMERS_LOCAL"], 400)
        self.assertEqual(months[1]["concept_totals"]["unclassified_out"], 25)
        self.assertEqual([row["running_balance"] for row in result["accounting_movements"]], [1400, 1195, 1170])
        self.assertNotIn("running_balance", data["accounting_movements"][0])

    def test_ledger_and_bank_views_never_add_the_same_movement_twice(self):
        data = accounting_fixture()
        data["bank_opening"] = 1000
        data["movements"] = [movement(50, "2024-01-05", 400)]
        result = build_snapshot(data)
        self.assertEqual(result["months"][0]["income"], 400)
        self.assertEqual(result["accounting_months"][0]["income"], 400)
        self.assertEqual(result["months"][0]["bank_end"], 1400)
        self.assertEqual(result["accounting_months"][0]["closing"], 1195)

    def test_accounting_sheets_survive_template_and_keep_source_distinction(self):
        snapshot = build_snapshot(accounting_fixture())
        template = template_workbook()
        template.create_sheet("Notas personales")["A1"] = "Mantener"
        for use_template in (None, saved(template)):
            with self.subTest(template=bool(use_template)):
                book = load_workbook(BytesIO(export_xlsx(snapshot, template=use_template)), data_only=True)
                self.assertEqual(book["Banco"]["G19"].value, "n.d.")
                self.assertFalse(any(cell.value is not None for row in book["Data"].iter_rows(min_row=9) for cell in row))
                summary = book["Resumen contable Odoo"]
                self.assertEqual([summary.cell(5, c).value for c in (2, 3, 4, 14)], [1000, 1195, 1170, 1000])
                self.assertEqual(summary["N6"].value, 400)
                self.assertEqual(summary["N7"].value, 230)
                self.assertEqual(summary["N8"].value, 1170)
                self.assertIsNone(summary["E5"].value)  # Months after cutoff.
                detail = book["Mayor bancario Odoo"]
                self.assertEqual(detail["I8"].value, 1170)
                self.assertEqual(detail["F6"].data_type, "s")
                self.assertEqual(detail["F6"].value, "=Not a formula")
                self.assertEqual(detail["N6"].value, "Sin transacción bancaria vinculada")
                self.assertIn("model=account.move.line", detail["Q6"].hyperlink.location)
                if use_template:
                    self.assertEqual(book["Notas personales"]["A1"].value, "Mantener")
                self.assertFalse(any(c.data_type == "e" for s in book for row in s for c in row))

    def test_monthly_deposit_heading_includes_unclassified_receipts_once(self):
        data = fixture()
        data["month"] = 2
        data["movements"].append(movement(90, "2024-02-02", 123, allocations=[]))
        for template in (None, saved(template_workbook())):
            output = export_xlsx(build_snapshot(data), template=template)
            values = load_workbook(BytesIO(output), data_only=True)["Banco"]
            formulas = load_workbook(BytesIO(output))["Banco"]
            self.assertEqual((values["G20"].value, values["I20"].value, values["AE20"].value), (400, 123, 523))
            self.assertIsNone(values["K20"].value)
            self.assertTrue(formulas["G20"].value.startswith("=SUM("))
            self.assertEqual(values["AE19"].value, 1000)
