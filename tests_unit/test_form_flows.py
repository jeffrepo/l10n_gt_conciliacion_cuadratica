from copy import deepcopy
from io import BytesIO
import unittest

from openpyxl import load_workbook

from core.calculation import build_snapshot
from core.export import export_xlsx
from core.flows import form_snapshot
from core.template import read_template
from test_accounting import accounting_fixture
from test_calculation import ledger, movement
from test_manual_balances import capture
from test_template import full_fixture, row_for, saved, template_workbook


def august_fixture():
    data = full_fixture()
    data.update(month=8, bank_opening=None, movements=[], controls={}, ledger=[])
    data["accounting_movements"] = []
    for ident, amount, code, partner in (
        (1, 600, "IN_CUSTOMERS_LOCAL", "Cliente local"),
        (2, 100, "IN_CUSTOMERS_FOREIGN", "Cliente extranjero"),
        (3, 200, "IN_RELATED", "Relacionada"),
        (4, 50, "", "Sin clasificar"),
        (5, -75, "OUT_SUPPLIERS", "Proveedor"),
    ):
        entry = ledger(ident, "2024-08-%02d" % ident, amount, company_amount=amount * 7.8)
        data["ledger"].append(entry)
        data["accounting_movements"].append({
            **entry, "partner": partner, "journal": "Depósitos" if amount > 0 else "Cheques",
            "reference": "=Referencia literal", "concept_code": code,
            "classification_origin": "rule" if code else "pending",
            "counterpart_accounts": "Clientes / proveedores", "statement_line_id": False,
        })
    data["manual_bank_balances"] = {"8": capture(1500000, 1500800)}
    return data


class TestFormFlows(unittest.TestCase):
    def test_august_concepts_detail_totals_and_independent_closing(self):
        for currency in ("GTQ", "USD"):
            data = august_fixture()
            data["metadata"]["currency"] = currency
            snapshot = build_snapshot(data)
            original = deepcopy(snapshot)
            for template in (None, saved(template_workbook())):
                with self.subTest(currency=currency, template=bool(template)):
                    raw = export_xlsx(snapshot, template=template)
                    book = load_workbook(BytesIO(raw), data_only=True)
                    formulas = load_workbook(BytesIO(raw))["Banco"]
                    main, detail = book["Banco"], book["Data"]
                    self.assertEqual(main["U19"].value, 1500000)
                    self.assertEqual(main["U20"].value, 950)
                    self.assertEqual([main.cell(row, 21).value for row in (21, 22, 23)], [600, 100, 200])
                    self.assertEqual(main["T16"].value, "Contabilidad Odoo")
                    self.assertGreaterEqual(main.row_dimensions[16].height, 28)
                    self.assertEqual(main["AE20"].value, 950)
                    self.assertEqual(main["G20"].value, 0)
                    self.assertIsNone(main["W20"].value)
                    self.assertEqual(main.cell(row_for(main, "Pendiente de clasificar"), 21).value, None)
                    self.assertEqual(main.cell(row_for(main, "Pendiente de clasificar"), 20).value, 50)
                    self.assertEqual(main.cell(row_for(main, "TOTAL EGRESOS"), 21).value, 75)
                    calculated = row_for(main, "SALDO CALCULADO ( A - B )")
                    reported = row_for(main, "Saldo final según banco (extracto o captura)")
                    adjusted = row_for(main, "SALDO CONCILIADO")
                    self.assertEqual(main.cell(calculated, 21).value, 1500875)
                    self.assertEqual(main.cell(reported, 21).value, 1500800)
                    self.assertEqual(main.cell(adjusted, 21).value, 1500800)
                    self.assertTrue(formulas.cell(adjusted, 21).value.startswith("=U%s+" % reported))
                    control = book["Control Odoo"] if template else main
                    self.assertEqual(control.cell(row_for(control, "Diferencia banco contra saldo informado"), 21).value, 75)
                    self.assertEqual(sum(detail.cell(row, 9).value or 0 for row in range(9, 14)), 950)
                    self.assertEqual(sum(detail.cell(row, 10).value or 0 for row in range(9, 14)), 75)
                    self.assertEqual(detail["K13"].value, 1500875)
                    self.assertEqual(detail["O9"].value, 4680)
                    self.assertEqual(detail["H9"].data_type, "s")
                    self.assertEqual(detail["X9"].value, "Contabilidad Odoo")
                    self.assertIsNone(detail["L9"].value)
                    self.assertEqual(detail["N11"].value, "Cuentas por cobrar a relacionadas locales y del exterior")
                    self.assertEqual(len([row for row in detail.iter_rows(min_row=9) if row[4].value]), 5)
                    self.assertFalse(any(cell.data_type == "e" for sheet in book for row in sheet for cell in row))
                    # A generated form is still a valid reusable template.
                    self.assertEqual(read_template(raw)["main"].title, "Banco")
            self.assertEqual(snapshot, original)

    def test_sources_are_chosen_monthly_and_never_added_together(self):
        data = accounting_fixture()
        data["bank_opening"] = 1000
        data["movements"] = [movement(50, "2024-01-05", 400)]
        snapshot = build_snapshot(data)
        view = form_snapshot(snapshot)
        self.assertEqual([month["flow_source"] for month in view["months"]], ["bank", "accounting", "accounting"])
        self.assertEqual([month["income"] for month in view["months"]], [400, 0, 0])
        self.assertEqual([month["expense"] for month in view["months"]], [0, 25, 0])
        self.assertEqual([row["source_model"] for row in view["movements"]], ["account.bank.statement.line", "account.move.line"])
        self.assertEqual([row["source_id"] for row in view["movements"]], [50, 13])
        self.assertIsNone(view["months"][1]["bank_adjusted"])

    def test_complete_zero_activity_statement_takes_priority(self):
        data = accounting_fixture()
        data["controls"] = {"1": {"source_id": 1, "name": "Enero sin actividad", "balance": 1000, "complete": True}}
        view = form_snapshot(build_snapshot(data))
        self.assertEqual(view["months"][0]["flow_source"], "bank")
        self.assertEqual(view["months"][0]["income"], 0)
        self.assertEqual(view["months"][1]["expense"], 25)

    def test_no_bank_or_accounting_evidence_remains_unknown(self):
        data = full_fixture()
        data.update(bank_opening=None, movements=[], controls={}, ledger=[])
        output = export_xlsx(build_snapshot(data))
        main = load_workbook(BytesIO(output), data_only=True)["Banco"]
        self.assertEqual(main["G20"].value, "n.d.")
        self.assertEqual(main["G21"].value, "n.d.")

    def test_captured_zero_does_not_turn_into_missing_opening(self):
        data = accounting_fixture()
        data["manual_bank_balances"] = {"1": capture(0, 195)}
        book = load_workbook(BytesIO(export_xlsx(build_snapshot(data))), data_only=True)
        main = book["Banco"]
        self.assertEqual(main["G19"].value, 0)
        self.assertEqual(main.cell(row_for(main, "SALDO CALCULADO ( A - B )"), 7).value, 195)
        self.assertEqual(main.cell(row_for(main, "SALDO CONCILIADO"), 7).value, 195)
