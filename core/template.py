"""Read and fill the supplied two-column-per-month XLSX form.

Templates supply presentation and manual text, never accounting balances.
The automated financial region and movement table are rebuilt on every export.
"""

from copy import copy
from io import BytesIO
import re
import unicodedata
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, BadZipFile, ZipFile

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.formula import Tokenizer
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table

from .calculation import MONTHS


def normalized(value):
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "", text)


ALIASES = {
    "IN_CUSTOMERS_LOCAL": ("Cuentas por cobrar clientes locales", "Clientes locales"),
    "IN_CUSTOMERS_FOREIGN": ("Cuentas por cobrar clientes del exterior", "Clientes del exterior"),
    "IN_RELATED": ("Cuentas y documentos por cobrar relacionadas", "Cuentas por cobrar a relacionadas"),
    "IN_SHAREHOLDERS": ("Cuentas por cobrar a socios",),
    "IN_EMPLOYEES": ("Cuentas por cobrar empleados", "Cuentas por cobrar a empleados"),
    "IN_ADVANCES": ("Anticipo de clientes", "Anticipos de clientes"),
    "IN_INTEREST": ("Intereses ganados",),
    "IN_TRANSFER": ("Trasferencia de fondos", "Transferencia de fondos", "Transferencias entre cuentas"),
    "other_in": ("Otros ingresos",),
    "IN_SUPPLIER_REFUNDS": ("Reintegro por devoluciones", "Reintegros por devoluciones"),
    "IN_LOAN_REPAYMENTS": ("Devoluciones por prestamos", "Devoluciones de prestamos"),
    "OUT_SUPPLIERS": ("Proveedores",),
    "OUT_OPERATING": ("Gastos Operativos",),
    "OUT_ADVANCES": ("Pagos anticipados",),
    "OUT_LOANS": ("Pago a prestamos", "Pago de prestamos"),
    "OUT_DIVIDENDS": ("Pago de dividendos",),
    "OUT_SHAREHOLDERS": ("Cuentas por pagar socios", "Cuentas por pagar a socios"),
    "OUT_RELATED_LOCAL": ("Cuentas por pagar relacionadas locales", "Cuentas por pagar a relacionadas locales"),
    "OUT_RELATED_FOREIGN": ("Cuentas por pagar relacionadas del exterior", "Cuentas por pagar a relacionadas del exterior"),
    "OUT_TRANSFER": ("Egreso por transferencia de fondos",),
    "other_out": ("Otros egresos",),
    "OUT_TAXES": ("Pago de impuestos",),
    "OUT_FEES": ("Comision ACH", "Comisiones bancarias"),
    "OUT_CHECKS": ("Cheques Emitidos",),
    "OUT_INTEREST": ("Intereses pagados",),
    "OUT_LOANS_GRANTED": ("Prestamos otorgados",),
    "OUT_REJECTED_DEPOSITS": ("Rechazos de depositos", "Depositos rechazados"),
    "bank_opening": ("Saldo inicial segun banco",),
    "available": ("Total depositos del periodo",),
    "expense": ("Total egresos",),
    "bank_end": ("Saldo final bancario",),
    "deposits": ("Depositos en transito",),
    "checks": ("Cheques en circulacion",),
    "book_balance": ("Saldo final segun libros",),
    "interest_pending": ("Ingresos por intereses ganados",),
    "debits_pending": ("Nota de debito no operadas", "Notas de debito no operadas"),
}
DETAILS = {
    "IN_RELATED": "IN_SHAREHOLDERS", "IN_SHAREHOLDERS": "IN_EMPLOYEES",
    "IN_TRANSFER": "other_in", "OUT_LOANS": "OUT_DIVIDENDS",
    "OUT_DIVIDENDS": "OUT_SHAREHOLDERS", "OUT_SHAREHOLDERS": "OUT_RELATED_LOCAL",
    "OUT_RELATED_LOCAL": "OUT_RELATED_FOREIGN", "OUT_RELATED_FOREIGN": "OUT_TRANSFER",
    "OUT_TRANSFER": "other_out",
}


def read_template(data):
    if len(data) > 8 * 1024 * 1024:
        raise ValueError("La plantilla XLSX no puede exceder 8 MB.")
    try:
        with ZipFile(BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 64 * 1024 * 1024:
                raise ValueError("La plantilla contiene demasiados datos. Use un formato sin históricos.")
            if any("vbaProject" in name or "externalLinks/" in name for name in archive.namelist()):
                raise ValueError("Use una plantilla XLSX sin macros ni enlaces a otros libros.")
        book = openpyxl.load_workbook(BytesIO(data), keep_links=False)
    except (BadZipFile, KeyError, OSError, ET.ParseError) as error:
        raise ValueError("El archivo no es una plantilla XLSX válida.") from error
    mains = [sheet for sheet in book if normalized(sheet["C19"].value) == "saldoinicialsegunbanco"]
    datas = [(sheet, row) for sheet in book for row in (7, 8)
             if normalized(sheet.cell(row, 1).value) == "mesdecobro"
             and normalized(sheet.cell(row, 9).value) == "ingreso"
             and normalized(sheet.cell(row, 10).value) == "egreso"]
    if len(mains) != 1 or len(datas) != 1 or mains[0] == datas[0][0]:
        raise ValueError("Use el formato de referencia: una hoja Banco con SALDO INICIAL SEGÚN BANCO en C19 y una hoja Data con encabezados en la fila 7 u 8.")
    main, (data_sheet, data_header) = mains[0], datas[0]
    if main.max_row > 2000 or main.max_column > 200 or data_sheet.max_row > 20000 or data_sheet.max_column > 200:
        raise ValueError("La plantilla excede el tamaño admitido. Retire históricos y formatos sobrantes.")
    for idx, month in enumerate(MONTHS):
        if normalized(main.cell(18, 6 + idx * 2).value) != normalized(month):
            raise ValueError("La fila 18 debe conservar enero a diciembre, con dos columnas por mes, desde F hasta AC.")
    if normalized(main["AD18"].value) not in ("total", "totales"):
        raise ValueError("La plantilla debe conservar TOTALES en AD18:AE18.")
    labels = {row: normalized(main.cell(row, 3).value) for row in range(19, main.max_row + 1)}
    row_map = {}
    for key, aliases in ALIASES.items():
        matches = [row for row, label in labels.items() if any(label.startswith(normalized(alias)) for alias in aliases)]
        if matches:
            row_map[key] = matches[0]
    totals = [row for row, label in labels.items() if label == "saldoconciliado"]
    if len(totals) != 2 or any(key not in row_map for key in ("available", "expense", "bank_end", "book_balance", "deposits", "checks")):
        raise ValueError("Faltan los totales o los dos SALDO CONCILIADO del formato de referencia.")
    row_map.update(bank_adjusted=totals[0], book_adjusted=totals[1], in_heading=20,
                   out_heading=row_map["available"] + 1, bank_adjustments=row_map["bank_end"] + 1,
                   book_adjustments=row_map["book_balance"] + 1)
    for key, start, end in (("other_bank", row_map["bank_end"], totals[0]), ("other_book", row_map["book_balance"], totals[1])):
        rows = [row for row, label in labels.items() if start < row < end and label.startswith("otros")]
        if rows:
            row_map[key] = rows[0]
    if "other_bank" in row_map:
        row_map["bank_note"] = row_map["other_bank"] + 1
    if "other_book" in row_map:
        row_map["receipts_pending"] = row_map["other_book"] + 1
        row_map["book_spacer"] = row_map["other_book"] + 2
    row_map["spacer"] = totals[0] + 1
    details = {}
    for key, next_key in DETAILS.items():
        if key not in row_map or next_key not in row_map:
            raise ValueError("Falta el bloque de desglose del concepto %s en la plantilla." % key)
        start, end = row_map[key] + 2, row_map[next_key]
        if start > end:
            raise ValueError("Los bloques de la plantilla no conservan el orden del formato de referencia.")
        bank_only = key in ("IN_TRANSFER", "OUT_TRANSFER")
        details[key] = [{"row": row, "label": tuple(str(main.cell(row, col).value or "") for col in ((3, 4) if bank_only else (3, 4, 5)))} for row in range(start, end)]
    return {"book": book, "main": main, "data": data_sheet, "data_header": data_header, "row_map": row_map,
            "details": details, "form_end": totals[1], "bank": str(main["D6"].value or ""),
            "account": str(main["D8"].value or "")}


def merge_manual_details(groups, slots, bank_only):
    """Match a manual label only to an unambiguous known account/name.

    Blank/missing identities never identify a bank. Unmatched manual labels
    remain with zero; unknown-source movements receive their own visible row.
    """
    remaining = list(groups)
    result = []
    for slot in slots:
        label = slot["label"]
        candidates = []
        account = normalized(label[-1])
        if account:
            candidates = [idx for idx, (key, _) in enumerate(remaining) if normalized(key[-1]) == account]
        elif label[0]:
            candidates = [idx for idx, (key, _) in enumerate(remaining) if normalized(key[0]) == normalized(label[0]) and not key[0].endswith("no identificado") and not key[0].endswith("no identificada")]
        if len(candidates) == 1:
            actual, values = remaining.pop(candidates[0])
            # Actual source names win; manual text supplies missing identity
            # only after matching an explicit account or unique known name.
            label = tuple(manual if (not real or "no identificad" in real.lower()) and manual else real for real, manual in zip(actual, label))
        else:
            values = {}
        result.append((label, values, slot["row"]))
    # Reuse empty template lines before inserting additional detail rows.
    for idx, (label, values, source_row) in enumerate(result):
        if not any(label) and remaining:
            actual, totals = remaining.pop(0)
            result[idx] = (actual, totals, source_row)
    fallback = slots[-1]["row"] if slots else None
    result.extend((key, values, fallback) for key, values in remaining)
    return result


def _copy_style(source, target):
    for attr in ("font", "fill", "border", "alignment", "protection"):
        setattr(target, attr, copy(getattr(source, attr)))
    target.number_format = source.number_format


def _copy_cell(source, target):
    if isinstance(source, MergedCell):
        return
    target.value = source.value
    target.data_type = source.data_type
    _copy_style(source, target)
    if source.comment:
        target.comment = copy(source.comment)
    if source.hyperlink:
        target.hyperlink = copy(source.hyperlink)


def _translate_formula(value, own_sheet, main_name, translate_row):
    if not isinstance(value, str) or not value.startswith("="):
        return value
    try:
        tokens = Tokenizer(value).items
        for token in tokens:
            if token.type != "OPERAND" or token.subtype != "RANGE":
                continue
            prefix, sep, cells = token.value.rpartition("!")
            if sep:
                if prefix.strip("'").replace("''", "'") != main_name:
                    continue
            elif own_sheet != main_name:
                continue
            if not sep:
                cells = token.value
            # Only actual A1 references; never alter named ranges or strings.
            if re.fullmatch(r"\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?", cells):
                cells = re.sub(r"(\$?[A-Z]{1,3}\$?)(\d+)", lambda match: match[1] + str(translate_row(int(match[2]))), cells)
                token.value = (prefix + sep if sep else "") + cells
        return "=" + "".join(token.value for token in tokens)
    except (ValueError, IndexError):
        return value


def fill_template(profile, generated, layout, snapshot):
    book, main, data_sheet = profile["book"], profile["main"], profile["data"]
    source_book = openpyxl.load_workbook(BytesIO(generated))
    source_values = openpyxl.load_workbook(BytesIO(generated), data_only=True)
    source_main = source_book["Banco"]
    old_end, new_end = profile["form_end"], layout["form_end"]
    originals = {(cell.row, cell.column): copy(cell) for row in main for cell in row}
    dimensions = {row: copy(value) for row, value in main.row_dimensions.items()}
    source_rows = {item["row"]: item.get("template_row") for item in layout["rows"]}
    backwards = {}
    for row, old in source_rows.items():
        if old:
            backwards.setdefault(old, row)
    # Anchor rows take precedence over prototypes reused by appended details.
    for key, old in profile["row_map"].items():
        if key in layout["index"]:
            backwards[old] = layout["index"][key]

    def translate_row(row):
        return row + new_end - old_end if row > old_end else backwards.get(row, row)

    original_merges = list(main.merged_cells.ranges)
    for merged in original_merges:
        owned_header = ((4 <= merged.min_row <= 12 and merged.min_col <= 7 and merged.max_col >= 2)
                        or (14 <= merged.min_row <= 18 and merged.min_col <= 31 and merged.max_col >= 2))
        if merged.max_row >= 19 or owned_header:
            main.unmerge_cells(str(merged))
    main.delete_rows(19, max(main.max_row - 18, 0))
    for row in list(main.row_dimensions):
        if row >= 19:
            del main.row_dimensions[row]
    for (row, col), source in originals.items():
        if row < 19 or isinstance(source, MergedCell):
            continue
        if row <= old_end and 2 <= col <= 31:
            continue
        target = main.cell(translate_row(row), col)
        _copy_cell(source, target)
        if source.data_type == "f":
            target.value = _translate_formula(source.value, main.title, main.title, translate_row)
    for row, dimension in dimensions.items():
        if row > old_end:
            new_row = translate_row(row)
            main.row_dimensions[new_row] = dimension
            main.row_dimensions[new_row].index = new_row
    for merged in original_merges:
        if merged.min_row > old_end or (merged.min_row >= 19 and merged.min_col > 31):
            main.merge_cells(start_row=translate_row(merged.min_row), end_row=translate_row(merged.max_row), start_column=merged.min_col, end_column=merged.max_col)
    cached = {}
    for row in source_main.iter_rows(min_row=4, max_row=new_end, min_col=2, max_col=31):
        for source in row:
            # Manual header notes outside the Odoo identity block are retained.
            if source.row <= 12 and source.column > 7:
                continue
            target = main.cell(source.row, source.column)
            _copy_cell(source, target)
            template_row = source.row if source.row < 19 else source_rows.get(source.row)
            prototype = originals.get((template_row, source.column))
            if prototype is not None and prototype.has_style:
                _copy_style(prototype, target)
                # Numeric/date types remain appropriate even if the template
                # used a text-style placeholder in an otherwise numeric slot.
                if source.data_type == "f" or isinstance(source.value, (int, float)) or source.number_format == "@":
                    target.number_format = source.number_format
                if source.row in (17, 18):
                    target.alignment = copy(source.alignment)
            if source.data_type == "f":
                cached[(main.title, source.coordinate)] = source_values["Banco"][source.coordinate].value
    for item in layout["rows"]:
        old_dimension = dimensions.get(item.get("template_row"))
        wanted = source_main.row_dimensions[item["row"]].height or 15
        main.row_dimensions[item["row"]].height = max(wanted, old_dimension.height or 15) if old_dimension else wanted
    for merged in source_main.merged_cells.ranges:
        if merged.max_row <= new_end:
            main.merge_cells(str(merged))
    # Refresh the small control panel present below the original form. Other
    # manual footer text and cells retain their values and styling.
    footer_keys = {"bancoinicial": "bank_opening", "containicial": "book_opening", "bancofinal": "statement_end", "contafinal": "book_balance", "diferenciabanco": "bank_difference", "diferenciacontabilidad": "difference"}
    for row in range(new_end + 1, main.max_row + 1):
        key = next((footer_keys[normalized(main.cell(row, col).value)] for col in range(2, 6) if normalized(main.cell(row, col).value) in footer_keys), None)
        if not key:
            continue
        for col in range(6, 32):
            main.cell(row, col).value = None
        for idx, month in enumerate(snapshot["months"]):
            main.cell(row, 7 + idx * 2, month.get(key) if month.get(key) is not None else "n.d.")
        month = snapshot["months"][0 if key.endswith("opening") else -1]
        main.cell(row, 31, month.get(key) if month.get(key) is not None else "n.d.")
    main.freeze_panes = "F19"
    main.print_title_rows = "17:18"
    main.print_title_cols = "B:E"
    main.print_area = "B4:AE%s" % main.max_row
    main.sheet_properties.pageSetUpPr.fitToPage = True
    main.page_setup.orientation = "landscape"
    main.page_setup.fitToWidth, main.page_setup.fitToHeight = 1, 0
    main.protection.sheet = False

    # Keep manual notes on other sheets and update direct references to form
    # totals when expanded detail blocks move those rows.
    for sheet in book:
        if sheet != main:
            for row in sheet:
                for cell in row:
                    if cell.data_type == "f":
                        cell.value = _translate_formula(cell.value, sheet.title, main.title, translate_row)
    for name in book.defined_names.values():
        if name.attr_text:
            name.attr_text = _translate_formula("=" + name.attr_text, "", main.title, translate_row)[1:]

    data_header = profile["data_header"]
    offset = data_header - 7
    original_data_styles = {(row, col): copy(data_sheet.cell(row + offset, col)) for row in (7, 8, 9) for col in range(1, 24)}
    table_names = list(data_sheet.tables)
    table_style = copy(data_sheet.tables[table_names[0]].tableStyleInfo) if table_names else None
    for name in table_names:
        del data_sheet.tables[name]
    for merged in list(data_sheet.merged_cells.ranges):
        if merged.max_row >= data_header and merged.min_col <= 23:
            data_sheet.unmerge_cells(str(merged))
    for row in data_sheet.iter_rows(min_row=data_header, max_col=23):
        for cell in row:
            cell.value = None
            cell.hyperlink = None
            cell.comment = None
    source_data = source_book["Data"]
    for row in source_data:
        for source in row:
            if source.row < 7 and source.coordinate not in ("D1", "D2", "D3"):
                continue
            target = data_sheet.cell(source.row + offset if source.row >= 7 else source.row, source.column)
            _copy_cell(source, target)
            prototype = original_data_styles.get((min(source.row, 9), source.column))
            if prototype is not None and prototype.has_style:
                _copy_style(prototype, target)
                target.number_format = source.number_format
    if not snapshot["metadata"]["bank"] and profile["bank"]:
        data_sheet["D3"] = profile["bank"]
    for col, dimension in source_data.column_dimensions.items():
        if dimension.min >= 12:
            data_sheet.column_dimensions[col] = copy(dimension)
    last_data_row = max(source_data.max_row, 8) + offset
    table = Table(displayName=table_names[0] if table_names else "MovimientosCQ", ref="A%s:K%s" % (data_header, last_data_row))
    table.tableStyleInfo = table_style
    data_sheet.add_table(table)
    data_sheet.auto_filter.ref = "A%s:W%s" % (data_header, last_data_row)
    data_sheet.freeze_panes = "A%s" % (data_header + 2)
    data_sheet.row_dimensions[data_header].height = max(data_sheet.row_dimensions[data_header].height or 15, 30)
    data_sheet.print_area = "A1:W%s" % last_data_row
    data_sheet.protection.sheet = False

    for name in ("Partidas conciliatorias", "Control Odoo"):
        if name in book.sheetnames:
            del book[name]
    pending = book.create_sheet("Partidas conciliatorias")
    for row in source_book["Partidas conciliatorias"]:
        for cell in row:
            _copy_cell(cell, pending.cell(cell.row, cell.column))
    for col, dimension in source_book["Partidas conciliatorias"].column_dimensions.items():
        pending.column_dimensions[col] = copy(dimension)
    pending.freeze_panes = "A2"
    pending.auto_filter.ref = source_book["Partidas conciliatorias"].auto_filter.ref
    control = book.create_sheet("Control Odoo")
    for row in source_main.iter_rows(min_row=new_end + 1):
        for source in row:
            _copy_cell(source, control.cell(source.row - new_end, source.column))
    for merged in source_main.merged_cells.ranges:
        if merged.min_row > new_end:
            control.merge_cells(start_row=merged.min_row - new_end, end_row=merged.max_row - new_end, start_column=merged.min_col, end_column=merged.max_col)
    for col, dimension in source_main.column_dimensions.items():
        control.column_dimensions[col] = copy(dimension)
    if (not snapshot["metadata"]["bank"] and profile["bank"]) or (not snapshot["metadata"]["account_number"] and profile["account"]):
        control.cell(control.max_row + 2, 3, "Los datos bancarios de cabecera que faltan en Odoo se conservaron de la plantilla. No identifican por sí solos los bancos de las contrapartes.")
    for sheet in (main, data_sheet, pending, control):
        sheet.sheet_view.showGridLines = False
    book.calculation.fullCalcOnLoad = True
    book.calculation.forceFullCalc = True
    output = BytesIO()
    book.save(output)
    return _formula_caches(output.getvalue(), book.sheetnames, cached)


def _formula_caches(data, sheet_names, cached):
    """Keep independently calculated values visible before Excel recalculates.

    openpyxl deliberately drops formula caches. Only module-owned formulas get
    caches from our calculation engine; custom formulas recalculate in Excel.
    """
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    output = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            contents = source.read(entry.filename)
            match = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", entry.filename)
            if match:
                sheet_name = sheet_names[int(match[1]) - 1]
                cells = {coordinate: value for (name, coordinate), value in cached.items() if name == sheet_name}
                if cells:
                    root = ET.fromstring(contents)
                    for cell in root.findall(".//{%s}c" % namespace):
                        value = cells.get(cell.get("r"))
                        if value is None or cell.find("{%s}f" % namespace) is None:
                            continue
                        cached_value = cell.find("{%s}v" % namespace)
                        if cached_value is None:
                            cached_value = ET.SubElement(cell, "{%s}v" % namespace)
                        cached_value.text = str(value)
                    contents = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(entry, contents)
    return output.getvalue()
