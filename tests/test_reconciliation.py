from odoo import Command
from odoo.addons.account.tests.common import AccountTestInvoicingCommon
from odoo.exceptions import AccessError, UserError, ValidationError
from odoo.tests import Form, tagged


@tagged("post_install", "-at_install")
class TestQuadraticReconciliation(AccountTestInvoicingCommon):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.company = cls.company_data["company"]
        cls.incoming = cls.env["account.account"].create({
            "name": "CQ pending receipts", "code": "CQ1001", "account_type": "asset_current",
            "reconcile": True, "company_ids": [Command.set(cls.company.ids)],
        })
        cls.outgoing = cls.env["account.account"].create({
            "name": "CQ pending payments", "code": "CQ1002", "account_type": "asset_current",
            "reconcile": True, "company_ids": [Command.set(cls.company.ids)],
        })
        cls.bank = cls.env["account.journal"].create({
            "name": "CQ test bank", "code": "CQBK", "type": "bank", "company_id": cls.company.id,
            "cq_enabled": True,
        })
        cls.bank.inbound_payment_method_line_ids.payment_account_id = cls.incoming
        cls.bank.outbound_payment_method_line_ids.payment_account_id = cls.outgoing
        cls.bank.bank_account_id = cls.env["res.partner.bank"].create({
            "acc_number": "TEST-CQ-0001", "partner_id": cls.company.partner_id.id,
            "company_id": cls.company.id,
        })
        cls.in_concept = cls.env.ref("l10n_gt_conciliacion_cuadratica.concept_in_customers_local")
        cls.out_concept = cls.env.ref("l10n_gt_conciliacion_cuadratica.concept_out_suppliers")

    def _statement(self, when, amount, balance_start, balance_end, cutoff, classified=True, journal=None):
        journal = journal or self.bank
        line = self.env["account.bank.statement.line"].create({
            "journal_id": journal.id, "date": when, "payment_ref": "CQ transaction",
            "amount": amount, "partner_id": self.partner_a.id,
            "cq_concept_id": (self.in_concept if amount >= 0 else self.out_concept).id if classified else False,
        })
        if line.move_id.state == "draft":
            line.move_id.action_post()
        statement = self.env["account.bank.statement"].create({
            "name": "CQ %s" % cutoff, "line_ids": [Command.link(line.id)],
            "balance_start": balance_start, "balance_end_real": balance_end, "cq_date_end": cutoff,
        })
        return statement, line

    def _report(self, month=1):
        return self.env["cq.report"]._generate(self.company, self.bank.default_account_id, 2024, month)

    def _related_journal(self, code, **values):
        journal = self.env["account.journal"].create({
            "name": "CQ %s" % code, "code": code, "type": "bank", "company_id": self.company.id,
            "default_account_id": self.bank.default_account_id.id,
            "suspense_account_id": self.bank.suspense_account_id.id,
            **values,
        })
        journal.inbound_payment_method_line_ids.payment_account_id = self.incoming
        journal.outbound_payment_method_line_ids.payment_account_id = self.outgoing
        return journal

    def test_real_statement_report_export_and_immutable_close(self):
        self._statement("2024-01-29", 100, 0, 100, "2024-01-31")
        report = self._report()
        self.assertEqual(report.issue_count, 0, report.payload["issues"])
        self.assertEqual(report.month_ids.bank_end, 100)
        self.assertEqual(report.month_ids.book_balance, 0)
        self.assertEqual(report.month_ids.book_adjustment, 100)
        self.assertEqual(report.month_ids.difference, 0)
        self.assertTrue(report._xlsx_bytes().startswith(b"PK"))
        report.action_export()
        self.assertTrue(report.file_data)
        report.action_confirm()
        self.assertEqual(report.state, "confirmed")
        with self.assertRaises(AccessError):
            report.write({"payload": {"forged": True}})
        with self.assertRaises(AccessError):
            report.with_context(cq_internal=True).write({"issue_count": 0})
        with self.assertRaises(AccessError):
            report.month_ids.write({"income": 999})

    def test_classification_changes_require_new_snapshot(self):
        _, line = self._statement("2024-01-29", 100, 0, 100, "2024-01-31", classified=False)
        old = self._report()
        self.assertEqual(old.unclassified_count, 1)
        with self.assertRaises(UserError):
            old.action_confirm()
        line.cq_concept_id = self.in_concept
        new = self._report()
        self.assertEqual(new.unclassified_count, 0)
        self.assertEqual(old.unclassified_count, 1)
        self.assertNotEqual(old.id, new.id)

    def test_rules_and_split_constraints(self):
        _, line = self._statement("2024-01-29", 100, 0, 100, "2024-01-31", classified=False)
        self.env["cq.rule"].create({
            "name": "Known customer", "company_id": self.company.id, "journal_id": self.bank.id,
            "concept_id": self.in_concept.id, "partner_ids": [Command.set(self.partner_a.ids)],
        })
        self.assertEqual(self._report().unclassified_count, 0)
        allocation = self.env["cq.allocation"].create({
            "statement_line_id": line.id, "concept_id": self.in_concept.id, "amount": 60,
        })
        self.assertEqual(self._report().unclassified_count, 1)
        with self.assertRaises(ValidationError), self.cr.savepoint():
            allocation.amount = 101
        with self.assertRaises(ValidationError), self.cr.savepoint():
            allocation.concept_id = self.out_concept

    def test_partial_payment_after_cutoff(self):
        self._statement("2024-01-10", 100, 0, 100, "2024-01-31")
        payment = self.env["account.payment"].create({
            "journal_id": self.bank.id, "date": "2024-01-20", "amount": 80,
            "payment_type": "outbound", "partner_type": "supplier", "partner_id": self.partner_a.id,
            "payment_method_line_id": self.bank.outbound_payment_method_line_ids[0].id,
        })
        payment.action_post()
        _, bank_line = self._statement("2024-02-05", -30, 100, 70, "2024-02-29")
        counterpart = bank_line.move_id.line_ids.filtered(lambda aml: aml.account_id == self.bank.suspense_account_id)
        counterpart.account_id = self.outgoing
        payment_line = payment.move_id.line_ids.filtered(lambda aml: aml.account_id == self.outgoing)
        (payment_line | counterpart).reconcile()
        january = self._report()
        february = self._report(2)
        self.assertEqual(january.month_ids.payments, 80)
        feb = february.month_ids.filtered(lambda item: item.number == 2)
        self.assertEqual(feb.payments, 50)
        self.assertEqual(feb.difference, 0)
        self.assertEqual(february.issue_count, 0, february.payload["issues"])

    def test_missing_coverage_cannot_be_confirmed(self):
        statement, _ = self._statement("2024-01-29", 100, 0, 100, "2024-01-31")
        statement.cq_date_end = False
        report = self._report()
        self.assertFalse(report.month_ids.control_available)
        self.assertIn("coverage", {issue["code"] for issue in report.payload["issues"]})
        with self.assertRaises(UserError):
            report.action_confirm()

    def test_foreign_bank_currency_is_preserved(self):
        foreign = self.setup_other_currency("EUR", rates=[("1900-01-01", 0.5)])
        self.bank.currency_id = foreign
        self.incoming.currency_id = foreign
        self.outgoing.currency_id = foreign
        self._statement("2024-01-29", 100, 0, 100, "2024-01-31")
        report = self._report()
        self.assertEqual(report.currency_id, foreign)
        self.assertEqual(report.month_ids.income, 100)
        self.assertEqual(report.month_ids.bank_end, 100)
        self.assertEqual(report.month_ids.difference, 0)

    def test_shared_suspense_does_not_mix_banks(self):
        self._statement("2024-01-29", 100, 0, 100, "2024-01-31")
        other = self.env["account.journal"].create({
            "name": "CQ other bank", "code": "CQB2", "type": "bank", "company_id": self.company.id,
            "suspense_account_id": self.bank.suspense_account_id.id,
        })
        self.env["account.bank.statement.line"].create({
            "journal_id": other.id, "date": "2024-01-20", "payment_ref": "Other bank", "amount": 900,
        })
        report = self._report()
        self.assertEqual(report.month_ids.ledger_suspense, -100)
        self.assertEqual(report.month_ids.income, 100)

    def test_explicit_origin_precedes_current_account_mapping(self):
        other = self.env["account.journal"].create({
            "name": "CQ source bank", "code": "CQS2", "type": "bank", "company_id": self.company.id,
        })
        source = self.env["account.bank.statement.line"].create({
            "journal_id": other.id, "date": "2024-01-20", "payment_ref": "Original bank", "amount": 25,
        })
        line = source.move_id.line_ids[:1]
        owner = self.env["cq.report"]._line_owner(line, {self.bank.default_account_id.id}, {})
        self.assertEqual(owner, other.default_account_id.id)

    def test_wizard_view_and_company_boundary(self):
        with Form(self.env["cq.generate.wizard"]) as form:
            form.company_id = self.company
            form.year = 2024
            form.month = "1"
            form.account_ids.add(self.bank.default_account_id)
        wizard = form.save()
        self._statement("2024-01-29", 100, 0, 100, "2024-01-31")
        wizard.action_calculate()
        self.assertEqual(len(wizard.report_ids), 1)
        self.assertEqual(wizard.state, "result")
        other_company = self.env["res.company"].create({"name": "CQ isolated company"})
        with self.assertRaises(AccessError):
            self.env["cq.report"].with_context(allowed_company_ids=self.company.ids)._generate(other_company, self.bank.default_account_id, 2024, 1)

    def test_three_journals_one_account_and_one_opening_balance(self):
        checks = self._related_journal("CQCH")
        transfers = self._related_journal("CQTR", cq_statement_source=True)
        self._statement("2024-01-10", 100, 0, 100, "2024-01-10")
        self._statement("2024-01-20", -20, 100, 80, "2024-01-20", journal=checks)
        self._statement("2024-01-25", -30, 80, 50, "2024-01-31", journal=transfers)
        # Archived / non-enabled journals remain part of the account's history.
        checks.active = False
        wizard = self.env["cq.generate.wizard"].create({
            "company_id": self.company.id, "year": 2024, "month": "1", "all_accounts": True,
        })
        wizard.action_calculate()
        self.assertEqual(len(wizard.report_ids), 1)
        report = wizard.report_ids
        self.assertEqual(report.account_id, self.bank.default_account_id)
        self.assertEqual(set(report.journal_ids.ids), set((self.bank | checks | transfers).ids))
        self.assertEqual(report.journal_id, transfers)
        self.assertEqual(report.month_ids.bank_opening, 0)
        self.assertEqual(report.month_ids.income, 100)
        self.assertEqual(report.month_ids.expense, 50)
        self.assertEqual(report.month_ids.bank_end, 50)
        self.assertEqual(report.month_ids.ledger_bank, 50)
        self.assertEqual(report.month_ids.ledger_suspense, -50)
        self.assertEqual(len(report.payload["movements"]), 3)
        self.assertEqual(report.issue_count, 0, report.payload["issues"])
        from io import BytesIO
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(report._xlsx_bytes()), data_only=True)
        self.assertEqual(workbook["Movimientos"].max_row, 4)
        self.assertEqual(workbook["Movimientos"]["U4"].value, transfers.display_name)
        self.assertIn(transfers.display_name, workbook["Conciliación"]["A9"].value)

    def test_payments_from_related_journals_and_other_bank_stay_separate(self):
        checks = self._related_journal("CQCH")
        transfers = self._related_journal("CQTR")
        self._statement("2024-01-10", 100, 0, 100, "2024-01-31")
        for journal, amount in [(checks, 20), (transfers, 30)]:
            payment = self.env["account.payment"].create({
                "journal_id": journal.id, "date": "2024-01-20", "amount": amount,
                "payment_type": "outbound", "partner_type": "supplier", "partner_id": self.partner_a.id,
                "payment_method_line_id": journal.outbound_payment_method_line_ids[0].id,
            })
            payment.action_post()
        other = self.env["account.journal"].create({
            "name": "Another ledger bank", "code": "CQB2", "type": "bank", "company_id": self.company.id,
            "cq_enabled": True, "suspense_account_id": self.bank.suspense_account_id.id,
        })
        other.outbound_payment_method_line_ids.payment_account_id = self.outgoing
        payment = self.env["account.payment"].create({
            "journal_id": other.id, "date": "2024-01-20", "amount": 900,
            "payment_type": "outbound", "partner_type": "supplier", "partner_id": self.partner_a.id,
            "payment_method_line_id": other.outbound_payment_method_line_ids[0].id,
        })
        payment.action_post()
        wizard = self.env["cq.generate.wizard"].create({
            "company_id": self.company.id, "year": 2024, "month": "1", "all_accounts": True,
        })
        wizard.action_calculate()
        self.assertEqual(len(wizard.report_ids), 2)
        report = wizard.report_ids.filtered(lambda item: item.account_id == self.bank.default_account_id)
        self.assertEqual(report.month_ids.payments, 50)
        self.assertEqual(report.month_ids.ledger_outstanding, -50)
        self.assertEqual(report.month_ids.bank_adjusted, 50)
        self.assertEqual(report.issue_count, 0, report.payload["issues"])
        other_report = wizard.report_ids - report
        self.assertEqual(other_report.month_ids.payments, 900)

    def test_multiple_statement_sources_require_explicit_control(self):
        other = self._related_journal("CQTR")
        self._statement("2024-01-10", 100, 0, 100, "2024-01-10")
        self._statement("2024-01-25", -30, 100, 70, "2024-01-31", journal=other)
        report = self._report()
        self.assertIn("statement_source", report.issue_ids.mapped("code"))
        self.assertIsNone(report.payload["months"][0]["bank_opening"])
        self.assertFalse(report.month_ids.control_available)
        self.assertEqual(report.month_ids.expense, 30)
        other.cq_statement_source = True
        self.assertEqual(self._report().issue_count, 0)
        with self.assertRaises(ValidationError), self.cr.savepoint():
            self.bank.cq_statement_source = True

    def test_rules_remain_scoped_to_source_journal(self):
        other = self._related_journal("CQTR", cq_statement_source=True)
        _, receipt = self._statement("2024-01-10", 100, 0, 100, "2024-01-10", classified=False)
        _, transfer = self._statement("2024-01-25", 30, 100, 130, "2024-01-31", classified=False, journal=other)
        other_concept = self.env.ref("l10n_gt_conciliacion_cuadratica.concept_in_interest")
        for journal, concept in [(self.bank, self.in_concept), (other, other_concept)]:
            self.env["cq.rule"].create({
                "name": journal.name, "company_id": self.company.id, "journal_id": journal.id,
                "concept_id": concept.id, "fallback": True,
            })
        report = self._report()
        self.assertEqual(report.issue_count, 0, report.payload["issues"])
        codes = {row["source_id"]: row["allocations"][0]["code"] for row in report.payload["movements"]}
        self.assertEqual(codes[receipt.id], self.in_concept.code)
        self.assertEqual(codes[transfer.id], other_concept.code)

    def test_bank_ledger_includes_general_journal_entries(self):
        self._statement("2024-01-10", 100, 0, 100, "2024-01-31")
        move = self.env["account.move"].create({
            "journal_id": self.company_data["default_journal_misc"].id, "date": "2024-01-20",
            "line_ids": [
                Command.create({"account_id": self.bank.default_account_id.id, "debit": 10}),
                Command.create({"account_id": self.company_data["default_account_revenue"].id, "credit": 10}),
            ],
        })
        move.action_post()
        report = self._report()
        self.assertEqual(report.month_ids.ledger_bank, 110)
        self.assertEqual(report.month_ids.difference, -10)
        self.assertIn("book_difference", report.issue_ids.mapped("code"))

    def test_group_rejects_mixed_currencies(self):
        foreign = self.setup_other_currency("EUR", rates=[("1900-01-01", 0.5)])
        self._related_journal("CQFX", currency_id=foreign.id)
        with self.assertRaises(UserError):
            self._report()

    def test_old_journal_snapshot_stays_readable(self):
        from ..models.report import INTERNAL
        self._statement("2024-01-10", 100, 0, 100, "2024-01-31")
        old = self._report()
        old.action_export()
        saved = old._xlsx_bytes()
        # Existing v1.0 records have only journal_id; update must not require
        # account_id, rewrite their payload, or regenerate their exported file.
        old.with_context(cq_internal=INTERNAL).write({"account_id": False, "journal_ids": [Command.clear()]})
        self.assertEqual(old._xlsx_bytes(), saved)
        action = old.action_new_version()
        self.assertEqual(action["context"]["default_account_ids"], [(6, 0, self.bank.default_account_id.ids)])

    def test_non_bank_and_wrong_allocation_are_rejected(self):
        with self.assertRaises(ValidationError), self.cr.savepoint():
            self.company_data["default_journal_misc"].cq_enabled = True
        _, line = self._statement("2024-01-29", 100, 0, 100, "2024-01-31")
        with self.assertRaises(ValidationError), self.cr.savepoint():
            line.cq_concept_id = self.out_concept
